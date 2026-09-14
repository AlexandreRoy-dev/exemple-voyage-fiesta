#!/usr/bin/env python3
"""Import Passagers.xlsx → GHL, owner = Yancie Auclair.

Creates missing contacts. Existing matches (email then phone) are reassigned
to Yancie; tags are appended, never replaced.
"""
from __future__ import annotations

import importlib.util
import json
import re
from datetime import date, datetime
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
XLSX = Path(r"c:\Users\Alex\AppData\Local\Temp\Passagers.xlsx")
OUT_DIR = ROOT / "imports" / "email-audience-extraction" / "batch-2026-09"
PAYLOADS = OUT_DIR / "yancie_auclair_payloads.json"
REPORT = OUT_DIR / "yancie_auclair_import_report.json"

YANCIE_USER_ID = "2k9aRsh2OkTdG7TnIiGo"
LOCATION_ID = "V90iyFBbBrCg3tpctRjc"
SOURCE = "Import Passagers Yancie Auclair 2026-09"
TAGS = ["import-conseiller", "import-yancie-auclair"]

spec = importlib.util.spec_from_file_location("deploy", HERE / "deploy-reservation-api-vps.py")
d = importlib.util.module_from_spec(spec)
spec.loader.exec_module(d)


def title_name(s: str) -> str:
    s = re.sub(r"\s+", " ", (s or "").strip())
    if not s:
        return ""
    parts = []
    for w in s.split(" "):
        if not w:
            continue
        low = w.lower()
        if low in {"de", "des", "du", "la", "le", "et"}:
            parts.append(low)
        else:
            parts.append(w[:1].upper() + w[1:].lower())
    return " ".join(parts)


def phone(s) -> str:
    digits = re.sub(r"\D+", "", str(s or ""))
    if digits.startswith("1") and len(digits) == 11:
        digits = digits[1:]
    if len(digits) != 10 or digits[0] == "0":
        return ""
    return "+1" + digits


def email(s) -> str:
    s = str(s or "").strip().lower()
    if "@" not in s:
        return ""
    host = s.split("@")[-1]
    if "." not in host or s.startswith("@") or s.endswith("@"):
        return ""
    return s


def clean_addr(s) -> str:
    s = str(s or "").strip()
    if s.lower() in {"", "x", "z", "xx", "n/a", "na", "-"}:
        return ""
    return s


def province(s) -> str:
    sl = str(s or "").strip().lower().replace(",", "").strip()
    if sl in {"", "ca", "canada"}:
        return ""
    if sl in {"quebec", "québec", "qc", "pq"}:
        return "QC"
    return str(s or "").strip()


def cell(headers, row, *names):
    for n in names:
        if n in headers:
            v = row[headers.index(n)]
            if v is None:
                return ""
            if isinstance(v, datetime):
                return v.date().isoformat()
            if isinstance(v, date):
                return v.isoformat()
            return str(v).strip()
    return ""


def header_matching(headers, *needles: str) -> str:
    for h in headers:
        low = h.lower()
        if all(n in low for n in needles):
            return h
    return ""


def parse_xlsx(path: Path) -> tuple[list[dict], dict]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    headers = [str(c).strip() if c is not None else "" for c in rows[0]]
    prenom = header_matching(headers, "pr") or headers[1]
    tel = header_matching(headers, "principal") or (headers[6] if len(headers) > 6 else "")
    cell_h = header_matching(headers, "cell")
    work_h = header_matching(headers, "travail")
    last_res = header_matching(headers, "servation")

    contacts = []
    skipped = 0
    for row in rows[1:]:
        last = title_name(cell(headers, row, "Nom"))
        first = title_name(cell(headers, row, prenom))
        em = email(cell(headers, row, "Courriel"))
        ph = (
            phone(cell(headers, row, tel))
            or phone(cell(headers, row, cell_h))
            or phone(cell(headers, row, work_h))
        )
        if not last and not first:
            skipped += 1
            continue
        if not em and not ph:
            skipped += 1
            continue
        dob = cell(headers, row, "Date naissance")
        if dob == "2001-01-01":
            dob = ""
        note_bits = []
        agent = cell(headers, row, "Agent")
        if agent:
            note_bits.append("Agent fichier: " + agent)
        last_r = cell(headers, row, last_res) if last_res else ""
        if last_r:
            note_bits.append("Dernière réservation: " + last_r)
        comments = cell(headers, row, "Commentaires")
        if comments:
            note_bits.append(comments)
        profil = cell(headers, row, "Profil")
        if profil:
            note_bits.append(profil[:1800])
        payload = {
            "firstName": first,
            "lastName": last,
            "email": em or None,
            "phone": ph or None,
            "address1": clean_addr(cell(headers, row, "Adresse")) or None,
            "city": title_name(clean_addr(cell(headers, row, "Ville"))) or None,
            "state": province(cell(headers, row, "Province")) or None,
            "postalCode": (clean_addr(cell(headers, row, "Postal")).upper().replace(" ", "") or None),
            "country": "CA",
            "dateOfBirth": dob or None,
            "source": SOURCE,
            "assignedTo": YANCIE_USER_ID,
            "tags": TAGS,
            "notes": "\n".join(note_bits) or None,
        }
        contacts.append({k: v for k, v in payload.items() if v not in (None, "")})

    seen: set[str] = set()
    uniq = []
    dups = 0
    for c in contacts:
        key = str(c.get("email") or c.get("phone"))
        if key in seen:
            dups += 1
            continue
        seen.add(key)
        uniq.append(c)

    stats = {
        "xlsx_rows": len(rows) - 1,
        "parsed": len(contacts),
        "unique": len(uniq),
        "skipped": skipped,
        "file_dups": dups,
        "with_email": sum(1 for c in uniq if c.get("email")),
        "with_phone": sum(1 for c in uniq if c.get("phone")),
    }
    return uniq, stats


REMOTE_JS = r"""
import { readFileSync, writeFileSync } from 'node:fs';

const env = {};
for (const line of readFileSync('/opt/voyage-fiesta-reservation/.env', 'utf8').split(/\r?\n/)) {
  const t = line.trim();
  if (!t || t.startsWith('#') || !t.includes('=')) continue;
  const i = t.indexOf('=');
  let v = t.slice(i + 1).trim();
  if ((v.startsWith('"') && v.endsWith('"')) || (v.startsWith("'") && v.endsWith("'"))) v = v.slice(1, -1);
  env[t.slice(0, i).trim()] = v;
}
const key = env.GHL_API_KEY;
const loc = env.GHL_LOCATION_ID || 'V90iyFBbBrCg3tpctRjc';
const ownerId = '2k9aRsh2OkTdG7TnIiGo';
const HDR = {
  Authorization: `Bearer ${key}`,
  Version: '2021-07-28',
  Accept: 'application/json',
  'Content-Type': 'application/json',
  'User-Agent': 'voyage-fiesta-import-yancie/1.0'
};
const contacts = JSON.parse(readFileSync('/tmp/yancie_import_payloads.json', 'utf8'));
const sleep = (ms) => new Promise((r) => setTimeout(r, ms));

async function ghl(path, opts = {}) {
  const res = await fetch(`https://services.leadconnectorhq.com${path}`, {
    ...opts,
    headers: { ...HDR, ...(opts.headers || {}) }
  });
  const text = await res.text();
  let data;
  try { data = JSON.parse(text); } catch { data = { raw: text.slice(0, 400) }; }
  return { ok: res.ok, status: res.status, data, text };
}

async function searchDuplicate(email, phone) {
  if (email) {
    const r = await ghl(`/contacts/search/duplicate?locationId=${encodeURIComponent(loc)}&email=${encodeURIComponent(email)}`);
    const id = r.data?.contact?.id || r.data?.id;
    if (id) return { id, assignedTo: r.data?.contact?.assignedTo || r.data?.assignedTo || '' };
  }
  if (phone) {
    const r = await ghl(`/contacts/search/duplicate?locationId=${encodeURIComponent(loc)}&number=${encodeURIComponent(phone)}`);
    const id = r.data?.contact?.id || r.data?.id;
    if (id) return { id, assignedTo: r.data?.contact?.assignedTo || r.data?.assignedTo || '' };
  }
  return null;
}

async function addTags(contactId, tags) {
  return ghl(`/contacts/${contactId}/tags`, {
    method: 'POST',
    body: JSON.stringify({ tags })
  });
}

async function addNote(contactId, body) {
  if (!body) return;
  return ghl(`/contacts/${contactId}/notes`, {
    method: 'POST',
    body: JSON.stringify({ body, userId: ownerId })
  });
}

async function createContact(payload) {
  const body = { ...payload, locationId: loc, assignedTo: ownerId };
  const notes = body.notes;
  delete body.notes;
  const r = await ghl('/contacts/', { method: 'POST', body: JSON.stringify(body) });
  if (!r.ok) {
    return { ok: false, status: r.status, error: r.data.message || r.data.msg || r.text.slice(0, 300) };
  }
  const contactId = r.data.contact?.id || r.data.id;
  if (contactId && notes) await addNote(contactId, notes);
  return { ok: true, contactId, status: r.status };
}

async function assignOwner(contactId, notes) {
  const r = await ghl(`/contacts/${contactId}`, {
    method: 'PUT',
    body: JSON.stringify({ assignedTo: ownerId })
  });
  if (!r.ok) {
    return { ok: false, status: r.status, error: r.data.message || r.data.msg || r.text.slice(0, 300) };
  }
  await addTags(contactId, ['import-conseiller', 'import-yancie-auclair']);
  if (notes) await addNote(contactId, notes);
  return { ok: true, contactId, status: r.status };
}

const report = { created: [], updated: [], skipped: [], errors: [] };
for (let i = 0; i < contacts.length; i++) {
  const p = contacts[i];
  const label = `${p.firstName || ''} ${p.lastName || ''} <${p.email || p.phone || ''}>`.trim();
  try {
    const existing = await searchDuplicate(p.email || '', p.phone || '');
    if (existing?.id) {
      const result = await assignOwner(existing.id, p.notes || '');
      if (result.ok) {
        report.updated.push({ label, contactId: existing.id, email: p.email, phone: p.phone });
        console.log('UPDATE', i + 1, label, existing.id);
      } else {
        report.errors.push({ label, status: result.status, error: result.error, email: p.email });
        console.log('ERR_UPDATE', i + 1, label, result.status, String(result.error || '').slice(0, 160));
      }
    } else {
      const result = await createContact(p);
      if (result.ok) {
        report.created.push({ label, contactId: result.contactId, email: p.email, phone: p.phone });
        console.log('CREATE', i + 1, label, result.contactId);
      } else {
        const err = String(result.error || '');
        if (result.status === 400 && /duplicate|exist/i.test(err)) {
          report.skipped.push({ label, reason: err, email: p.email, phone: p.phone });
          console.log('SKIP_DUP', i + 1, label, err.slice(0, 120));
        } else {
          report.errors.push({ label, status: result.status, error: err, email: p.email });
          console.log('ERR', i + 1, label, result.status, err.slice(0, 200));
        }
      }
    }
  } catch (e) {
    report.errors.push({ label, error: String(e) });
    console.log('EXC', i + 1, label, e);
  }
  await sleep(220);
}

report.summary = {
  total: contacts.length,
  created: report.created.length,
  updated: report.updated.length,
  skipped: report.skipped.length,
  errors: report.errors.length,
  assignedTo: ownerId
};
writeFileSync('/tmp/yancie_import_report.json', JSON.stringify(report, null, 2));
console.log('SUMMARY', JSON.stringify(report.summary));
"""


def main() -> None:
    if not XLSX.exists():
        raise SystemExit(f"xlsx missing: {XLSX}")
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    payloads, stats = parse_xlsx(XLSX)
    PAYLOADS.write_text(json.dumps(payloads, ensure_ascii=False, indent=2), encoding="utf-8")
    print("PARSE", json.dumps(stats, ensure_ascii=False))

    c = d.connect()
    try:
        sftp = c.open_sftp()
        sftp.put(str(PAYLOADS), "/tmp/yancie_import_payloads.json")
        with sftp.file("/tmp/yancie_import.mjs", "w") as f:
            f.write(REMOTE_JS)
        sftp.close()
        code, out, err = d.run(c, "node /tmp/yancie_import.mjs", timeout=2700)
        if out:
            print(out[-6000:] if len(out) > 6000 else out)
        if err:
            print(err[-1500:])
        sftp = c.open_sftp()
        sftp.get("/tmp/yancie_import_report.json", str(REPORT))
        sftp.close()
        report = json.loads(REPORT.read_text(encoding="utf-8"))
        print("REPORT", report.get("summary"))
        print("exit", code)
    finally:
        c.close()


if __name__ == "__main__":
    main()
