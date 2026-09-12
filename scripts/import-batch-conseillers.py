#!/usr/bin/env python3
"""Batch import Passagers xlsx → GHL create-only, assigned to conseiller.

Skips contacts that already exist (email or phone). Does not update existing.
"""
from __future__ import annotations

import importlib.util
import json
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
BATCH = ROOT / "imports" / "email-audience-extraction" / "batch-2026-09"
LOCATION_ID = "V90iyFBbBrCg3tpctRjc"

# Remaining agents — users already exist; never create GHL users.
AGENTS = [
    {
        "first": "Fannie",
        "last": "Hovington",
        "email": "fannie@voyagefiesta.com",
        "xlsx": BATCH / "FANNIE_HOVINGTON__Passagers (37).xlsx",
        "tag": "import-fannie-hovington",
    },
    {
        "first": "Fanny",
        "last": "Veillette",
        "email": "fanny@voyagefiesta.com",
        "xlsx": BATCH / "FANNY_VEILLETTE__Passagers (38).xlsx",
        "tag": "import-fanny-veillette",
    },
    {
        "first": "Geneviève",
        "last": "Gendron",
        "email": "genevieve.voyagefiesta@gmail.com",
        "xlsx": BATCH / "GENEVIEVE_GENDRON__Passagers (39).xlsx",
        "tag": "import-genevieve-gendron",
    },
    {
        "first": "Ikram",
        "last": "Khlamech",
        "email": "ikram.voyagefiesta@gmail.com",
        "xlsx": BATCH / "Ikram_Khlamech__Passagers (40).xlsx",
        "tag": "import-ikram-khlamech",
    },
    {
        "first": "Isabelle",
        "last": "Couture",
        "email": "isabellec@voyagefiesta.com",
        "xlsx": BATCH / "Isabelle_Couture__Passagers (41).xlsx",
        "tag": "import-isabelle-couture",
    },
    {
        "first": "Jody",
        "last": "Foisy",
        "email": "jody.f@voyagefiesta.com",
        "xlsx": BATCH / "Jody_Foisy__Passagers (42).xlsx",
        "tag": "import-jody-foisy",
    },
    {
        "first": "Karine",
        "last": "Tremblay",
        "email": "karine.tremblay@voyagefiesta.com",
        "xlsx": BATCH / "KARINE_TREMBLAY__Passagers (43).xlsx",
        "tag": "import-karine-tremblay",
    },
    {
        "first": "Kevin",
        "last": "Watier",
        "email": "kevin@voyagefiesta.com",
        "xlsx": BATCH / "KEVIN_WATIER__Passagers (44).xlsx",
        "tag": "import-kevin-watier",
    },
    {
        "first": "Liette",
        "last": "Grondin",
        "email": "liette@experiencesvip.ca",
        "xlsx": BATCH / "LIETTE_GRONDIN_AUDIENCE__Passagers (45).xlsx",
        "tag": "import-liette-grondin",
    },
    {
        "first": "Lisanne",
        "last": "Turcotte",
        "email": "lisanne@voyagefiesta.com",
        "xlsx": BATCH / "LISANNE_TURCOTTE__Passagers (46).xlsx",
        "tag": "import-lisanne-turcotte",
    },
    {
        "first": "Marie-Ève",
        "last": "Dubois",
        "email": "marie-eved@voyagefiesta.com",
        "xlsx": next(BATCH.glob("*Dubois*Passagers*.xlsx"), None),
        "tag": "import-marie-eve-dubois",
    },
    {
        "first": "Mathieu",
        "last": "Ducharme",
        "email": "travel.avec.math@gmail.com",
        "xlsx": BATCH / "MATHIEU_DUCHARME__Passagers (48).xlsx",
        "tag": "import-mathieu-ducharme",
    },
    {
        "first": "Monika",
        "last": "Fortin",
        "email": "monika@voyagefiesta.com",
        "xlsx": BATCH / "MONIKA_FORTIN__Passagers (49).xlsx",
        "tag": "import-monika-fortin",
    },
    {
        "first": "Nicolas",
        "last": "Belley",
        "email": "nicolas.b@voyagefiesta.com",
        "xlsx": BATCH / "Nicolas_Belley__Passagers (50).xlsx",
        "tag": "import-nicolas-belley",
    },
    {
        "first": "Nicolas",
        "last": "Rodrigue",
        "email": "nicolas.rodrigue@voyagefiesta.com",
        "xlsx": BATCH / "Nicolas_Rodrigue__Passagers (51).xlsx",
        "tag": "import-nicolas-rodrigue",
    },
    {
        "first": "Noémie",
        "last": "Vincent",
        "email": "noemie@voyagefiesta.com",
        "xlsx": next(BATCH.glob("*Vincent*Passagers*.xlsx"), None),
        "tag": "import-noemie-vincent",
    },
    {
        "first": "Odrey",
        "last": "Beaulieu",
        "email": "odrey.beaulieu@voyagefiesta.com",
        "xlsx": BATCH / "Odrey_Beaulieu__Passagers (53).xlsx",
        "tag": "import-odrey-beaulieu",
    },
    {
        "first": "Pascale",
        "last": "El Chemali",
        "email": "pascale@voyagefiesta.com",
        "xlsx": BATCH / "Pascale_El_Chemali__Passagers (55).xlsx",
        "tag": "import-pascale-el-chemali",
    },
    {
        "first": "Patrick",
        "last": "O'Farrell",
        "email": "patrick@voyagefiesta.ca",
        "xlsx": BATCH / "Patrick_O_Farrell__Passagers (56).xlsx",
        "tag": "import-patrick-ofarrell",
    },
    {
        "first": "Sophie",
        "last": "Denoncourt",
        "email": "sophie.voyagefiesta@gmail.com",
        "xlsx": BATCH / "Sophie_Denoncourt__Passagers (54).xlsx",
        "tag": "import-sophie-denoncourt",
    },
]

spec = importlib.util.spec_from_file_location(
    "deploy", HERE / "deploy-reservation-api-vps.py"
)
d = importlib.util.module_from_spec(spec)
spec.loader.exec_module(d)


def norm_key(k: str) -> str:
    return (
        k.lower()
        .replace("é", "e")
        .replace("è", "e")
        .replace("ê", "e")
        .replace("à", "a")
        .replace("ô", "o")
        .strip()
    )


def pick(row: dict, *names: str) -> str:
    want = {norm_key(n) for n in names}
    for k, v in row.items():
        nk = norm_key(k)
        nk_ascii = re.sub(r"[^a-z0-9. ()#]+", "", nk)
        if nk in want or nk_ascii in want or any(n in nk or n in nk_ascii for n in want):
            if v is None:
                continue
            s = str(v).strip()
            if s and s.lower() not in ("none", "null"):
                return s
    return ""


def pick_contains(row: dict, *needles: str) -> str:
    for k, v in row.items():
        nk = norm_key(k)
        nk_ascii = re.sub(r"[^a-z0-9. ()#]+", "", nk)
        if all(n in nk or n in nk_ascii for n in needles):
            if v is None:
                continue
            s = str(v).strip()
            if s and s.lower() not in ("none", "null"):
                return s
    return ""


def normalize_phone(raw: str) -> str:
    digits = re.sub(r"\D", "", raw or "")
    if not digits:
        return ""
    if len(digits) == 10:
        return f"+1{digits}"
    if len(digits) == 11 and digits.startswith("1"):
        return f"+{digits}"
    if (raw or "").strip().startswith("+"):
        return "+" + digits
    return f"+{digits}" if digits else ""


def normalize_email(raw: str) -> str:
    return str(raw or "").strip().lower()


def read_xlsx(path: Path) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    headers = [str(c).strip() if c is not None else "" for c in rows[0]]
    out = []
    for r in rows[1:]:
        if not any(c is not None and str(c).strip() for c in r):
            continue
        out.append(
            {
                headers[i]: ("" if r[i] is None else str(r[i]).strip())
                for i in range(min(len(headers), len(r)))
            }
        )
    return out


def build_payloads(rows: list[dict], user_id: str, source: str, tag: str) -> list[dict]:
    out = []
    for row in rows:
        last = ""
        first = ""
        for k, v in row.items():
            nk = re.sub(r"[^a-z]", "", norm_key(k))
            if nk == "nom":
                last = str(v or "").strip()
            elif nk in ("prenom", "prnom") or (
                nk.startswith("pr") and nk.endswith("nom") and nk != "nom"
            ):
                first = str(v or "").strip()
        email = normalize_email(pick(row, "courriel", "email"))
        p1 = normalize_phone(pick_contains(row, "principal"))
        p2 = normalize_phone(pick_contains(row, "cell"))
        phone = p1 or p2
        if not email and not phone:
            continue
        address = pick(row, "adresse")
        city = pick(row, "ville")
        state = pick(row, "province")
        postal = pick(row, "postal")
        comments = pick(row, "commentaires")
        profil = pick(row, "profil")
        notes_bits = [b for b in (comments, profil) if b]
        payload = {
            "firstName": first.title() if first else "",
            "lastName": last.title() if last else "",
            "email": email or None,
            "phone": phone or None,
            "address1": address or None,
            "city": city.title() if city else None,
            "state": state or None,
            "postalCode": postal.replace(" ", "") if postal else None,
            "country": "CA",
            "source": source,
            "assignedTo": user_id,
            "tags": ["import-conseiller", tag],
            "notes": "\n".join(notes_bits) if notes_bits else None,
        }
        payload = {k: v for k, v in payload.items() if v not in (None, "")}
        out.append(payload)
    return out


REMOTE_JS = r"""
import { readFileSync, writeFileSync } from 'node:fs';

const env = {};
for (const line of readFileSync('/opt/voyage-fiesta-reservation/.env', 'utf8').split(/\r?\n/)) {
  const t = line.trim();
  if (!t || t.startsWith('#') || !t.includes('=')) continue;
  const i = t.indexOf('=');
  let v = t.slice(i + 1).trim();
  if ((v.startsWith('"') && v.endsWith('"')) || (v.startsWith("'") && v.endsWith("'"))) v = v.slice(1, -1);
  env[t.slice(0, i).trim()] = v;
}
const key = env.GHL_API_KEY;
const loc = env.GHL_LOCATION_ID || 'V90iyFBbBrCg3tpctRjc';
const HDR = {
  Authorization: `Bearer ${key}`,
  Version: '2021-07-28',
  Accept: 'application/json',
  'Content-Type': 'application/json',
  'User-Agent': 'voyage-fiesta-import/1.0'
};
const job = JSON.parse(readFileSync('/tmp/batch_import_job.json', 'utf8'));
const sleep = (ms) => new Promise((r) => setTimeout(r, ms));

async function listUsers() {
  const res = await fetch(`https://services.leadconnectorhq.com/users/?locationId=${loc}`, { headers: HDR });
  const data = await res.json();
  return data.users || [];
}

async function createUser(agent) {
  const body = {
    companyId: env.GHL_COMPANY_ID || undefined,
    firstName: agent.first,
    lastName: agent.last,
    email: agent.email,
    password: `Vf!${Math.random().toString(36).slice(2)}9A`,
    type: 'account',
    role: 'user',
    locationIds: [loc]
  };
  // strip undefined
  Object.keys(body).forEach((k) => body[k] === undefined && delete body[k]);
  const res = await fetch('https://services.leadconnectorhq.com/users/', {
    method: 'POST',
    headers: HDR,
    body: JSON.stringify(body)
  });
  const text = await res.text();
  let data; try { data = JSON.parse(text); } catch { data = { raw: text }; }
  return { ok: res.ok, status: res.status, data, text: text.slice(0, 400) };
}

async function searchDuplicate(email, phone) {
  if (email) {
    const res = await fetch('https://services.leadconnectorhq.com/contacts/search', {
      method: 'POST', headers: HDR,
      body: JSON.stringify({ locationId: loc, pageLimit: 5, filters: [{ field: 'email', operator: 'eq', value: email }] })
    });
    const data = await res.json();
    if ((data.contacts || [])[0]) return data.contacts[0];
  }
  if (phone) {
    const res = await fetch('https://services.leadconnectorhq.com/contacts/search', {
      method: 'POST', headers: HDR,
      body: JSON.stringify({ locationId: loc, pageLimit: 5, filters: [{ field: 'phone', operator: 'eq', value: phone }] })
    });
    const data = await res.json();
    if ((data.contacts || [])[0]) return data.contacts[0];
  }
  return null;
}

async function createContact(payload) {
  const body = { ...payload, locationId: loc };
  const notes = body.notes;
  delete body.notes;
  const res = await fetch('https://services.leadconnectorhq.com/contacts/', {
    method: 'POST', headers: HDR, body: JSON.stringify(body)
  });
  const text = await res.text();
  let data; try { data = JSON.parse(text); } catch { data = { raw: text }; }
  if (!res.ok) return { ok: false, status: res.status, error: data.message || data.msg || text.slice(0, 300) };
  const contactId = data.contact?.id || data.id;
  if (contactId && notes) {
    await fetch(`https://services.leadconnectorhq.com/contacts/${contactId}/notes`, {
      method: 'POST', headers: HDR, body: JSON.stringify({ body: notes })
    });
  }
  return { ok: true, contactId, status: res.status };
}

const users = await listUsers();
const byEmail = new Map(users.map((u) => [String(u.email || '').toLowerCase(), u]));
const report = { agents: [] };

for (const agent of job.agents) {
  const email = String(agent.email || '').toLowerCase();
  let user = byEmail.get(email);
  const agentReport = {
    name: `${agent.first} ${agent.last}`,
    email,
    userId: null,
    userCreated: false,
    userError: null,
    total: (agent.contacts || []).length,
    created: [],
    skipped: [],
    errors: []
  };

  if (!user) {
    agentReport.userError = 'existing user required — skipped (no user create)';
    console.log('USER_MISSING_SKIP', email);
    report.agents.push(agentReport);
    continue;
  }

  const userId = user.id;
  agentReport.userId = userId;
  console.log('AGENT', agentReport.name, userId, 'contacts', agent.contacts.length);

  for (let i = 0; i < agent.contacts.length; i++) {
    const p = { ...agent.contacts[i], assignedTo: userId };
    const label = `${p.firstName || ''} ${p.lastName || ''} <${p.email || p.phone || ''}>`.trim();
    try {
      const existing = await searchDuplicate(p.email || '', p.phone || '');
      if (existing?.id) {
        agentReport.skipped.push({ label, contactId: existing.id });
        console.log('SKIP', agent.last, i + 1, label, existing.id);
      } else {
        const result = await createContact(p);
        if (result.ok) {
          agentReport.created.push({ label, contactId: result.contactId });
          console.log('CREATE', agent.last, i + 1, label, result.contactId);
        } else {
          const err = String(result.error || '');
          if (result.status === 400 && /duplicate|exist/i.test(err)) {
            agentReport.skipped.push({ label, reason: err });
            console.log('SKIP_DUP', agent.last, i + 1, label);
          } else {
            agentReport.errors.push({ label, status: result.status, error: err });
            console.log('ERR', agent.last, i + 1, label, result.status, err.slice(0, 160));
          }
        }
      }
    } catch (e) {
      agentReport.errors.push({ label, error: String(e) });
      console.log('EXC', agent.last, i + 1, e);
    }
    await sleep(220);
  }

  agentReport.summary = {
    created: agentReport.created.length,
    skipped: agentReport.skipped.length,
    errors: agentReport.errors.length
  };
  report.agents.push(agentReport);
}

report.totals = {
  agents: report.agents.length,
  created: report.agents.reduce((s, a) => s + a.created.length, 0),
  skipped: report.agents.reduce((s, a) => s + a.skipped.length, 0),
  errors: report.agents.reduce((s, a) => s + a.errors.length, 0)
};
writeFileSync('/tmp/batch_import_report.json', JSON.stringify(report, null, 2));
console.log('TOTALS', JSON.stringify(report.totals));
for (const a of report.agents) {
  console.log('SUMMARY_AGENT', a.name, JSON.stringify({ userId: a.userId, userCreated: a.userCreated, userError: a.userError, ...a.summary }));
}
"""


def main() -> None:
    job_agents = []
    for agent in AGENTS:
        xlsx = agent["xlsx"]
        if xlsx is None or not Path(xlsx).exists():
            print(f"SKIP_NO_XLSX {agent['first']} {agent['last']}")
            continue
        rows = read_xlsx(Path(xlsx))
        source = f"Import {agent['first']} {agent['last']} 2026-09"
        # user id filled on server; placeholder for tags/source
        payloads = build_payloads(rows, "PENDING", source, agent["tag"])
        for p in payloads:
            p.pop("assignedTo", None)
        job_agents.append(
            {
                "first": agent["first"],
                "last": agent["last"],
                "email": agent["email"],
                "tag": agent["tag"],
                "contacts": payloads,
            }
        )
        print(f"READY {agent['first']} {agent['last']}: {len(payloads)} contacts from {Path(xlsx).name}")

    job = {"agents": job_agents}
    local_job = BATCH / "batch_import_job_remaining.json"
    local_job.write_text(json.dumps(job, ensure_ascii=False, indent=2), encoding="utf-8")
    report_path = BATCH / "batch_import_report_remaining.json"

    c = d.connect()
    try:
        sftp = c.open_sftp()
        sftp.put(str(local_job), "/tmp/batch_import_job.json")
        with sftp.file("/tmp/batch_import.mjs", "w") as f:
            f.write(REMOTE_JS)
        sftp.close()
        # ~1.3s/contact; Fanny alone is ~430 rows → allow 45 min
        code, out, err = d.run(c, "node /tmp/batch_import.mjs", timeout=2700)
        print(out[-6000:] if len(out) > 6000 else out)
        if err:
            print(err[-1500:])
        sftp = c.open_sftp()
        sftp.get("/tmp/batch_import_report.json", str(report_path))
        sftp.close()
        report = json.loads(report_path.read_text(encoding="utf-8"))
        print("TOTALS", report.get("totals"))
        for a in report.get("agents", []):
            print(
                a["name"],
                a.get("userId"),
                a.get("summary"),
                a.get("userError") or "",
            )
    finally:
        c.close()


if __name__ == "__main__":
    main()
